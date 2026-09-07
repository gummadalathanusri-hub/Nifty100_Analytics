from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.screener.engine import ScreenerEngine


OUTPUT_PATH = Path("output/screener_output.xlsx")

PRESETS = {
    "quality_compounder": {
        "roe_min": 15,
        "de_max": 1.0,
        "fcf_min": 0,
        "revenue_cagr_5yr_min": 10,
    },
    "value_pick": {
        "pe_max": 20,
        "pb_max": 3.0,
        "de_max": 2.0,
        "dividend_yield_min": 1,
    },
    "growth_accelerator": {
        "pat_cagr_5yr_min": 20,
        "revenue_cagr_5yr_min": 15,
        "de_max": 2.0,
    },
    "dividend_champion": {
        "dividend_yield_min": 2,
        "dividend_payout_max": 80,
        "fcf_min": 0,
    },
    "debt_free_blue_chip": {
        "de_max": 0,
        "de_exact": True,
        "roe_min": 12,
        "sales_min": 5000,
    },
    "turnaround_watch": {
        "revenue_cagr_3yr_min": 10,
        "fcf_min": 0,
        "de_declining": True,
    },
}


KPI_COLUMNS = [
    "company_id",
    "year",
    "sales_cr",
    "net_profit_cr",
    "return_on_equity_pct",
    "roce_pct",
    "net_profit_margin_pct",
    "operating_profit_margin_pct",
    "debt_to_equity",
    "interest_coverage",
    "free_cash_flow_cr",
    "fcf_cagr",
    "cash_from_operations_cr",
    "cfo_pat_ratio",
    "revenue_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",
    "asset_turnover",
    "pe_ratio",
    "pb_ratio",
    "dividend_yield_pct",
    "composite_quality_score",
]


def apply_formatting(path: Path) -> None:
    workbook = load_workbook(path)

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )
    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                28,
            )

        headers = {
            cell.value: cell.column
            for cell in worksheet[1]
        }

        threshold_rules = {
            "return_on_equity_pct": ("greaterThan", 15),
            "debt_to_equity": ("lessThan", 1),
            "free_cash_flow_cr": ("greaterThanOrEqual", 0),
            "revenue_cagr_5yr": ("greaterThan", 10),
            "pat_cagr_5yr": ("greaterThan", 20),
            "pe_ratio": ("lessThan", 20),
            "pb_ratio": ("lessThan", 3),
            "dividend_yield_pct": ("greaterThan", 1),
        }

        for column_name, (operator, value) in threshold_rules.items():
            if column_name not in headers:
                continue

            column_letter = get_column_letter(headers[column_name])
            cell_range = f"{column_letter}2:{column_letter}{worksheet.max_row}"

            if operator == "greaterThan":
                formula = f"{column_letter}2>{value}"
                rule = CellIsRule(
                    operator="greaterThan",
                    formula=[str(value)],
                    fill=green_fill,
                )
            elif operator == "greaterThanOrEqual":
                rule = CellIsRule(
                    operator="greaterThanOrEqual",
                    formula=[str(value)],
                    fill=green_fill,
                )
            else:
                rule = CellIsRule(
                    operator="lessThan",
                    formula=[str(value)],
                    fill=green_fill,
                )

            worksheet.conditional_formatting.add(
                cell_range,
                rule,
            )

        for row in worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
        ):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"

    workbook.save(path)


def build_export() -> None:
    engine = ScreenerEngine()

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        OUTPUT_PATH,
        engine="openpyxl",
    ) as writer:
        for preset_name, filters in PRESETS.items():
            result = engine.run(filters)

            columns = [
                column
                for column in KPI_COLUMNS
                if column in result.columns
            ]

            export_df = result[columns].copy()

            if "composite_quality_score" in export_df.columns:
                export_df = export_df.sort_values(
                    "composite_quality_score",
                    ascending=False,
                    na_position="last",
                )

            sheet_name = preset_name[:31]
            export_df.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

    apply_formatting(OUTPUT_PATH)

    print(f"Created: {OUTPUT_PATH}")
    print(f"Sheets: {len(PRESETS)}")


if __name__ == "__main__":
    build_export()