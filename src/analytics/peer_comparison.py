from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DB_PATH = Path("nifty100.db")
PEER_GROUPS_PATH = Path("data/supporting/peer_groups.xlsx")
OUTPUT_PATH = Path("output/peer_comparison.xlsx")

METRICS = [
    "ROE",
    "ROCE",
    "Net Profit Margin",
    "D/E",
    "FCF",
    "PAT CAGR 5yr",
    "Revenue CAGR 5yr",
    "EPS CAGR 5yr",
    "Interest Coverage",
    "Asset Turnover",
]

METRIC_SOURCE = {
    "ROE": "return_on_equity_pct",
    "ROCE": "roce_pct",
    "Net Profit Margin": "net_profit_margin_pct",
    "D/E": "debt_to_equity",
    "FCF": "free_cash_flow_cr",
    "PAT CAGR 5yr": "pat_cagr_5yr",
    "Revenue CAGR 5yr": "revenue_cagr_5yr",
    "EPS CAGR 5yr": "eps_cagr_5yr",
    "Interest Coverage": "interest_coverage",
    "Asset Turnover": "asset_turnover",
}

def load_data():
    conn = __import__("sqlite3").connect(DB_PATH)

    ratios = pd.read_sql_query(
        "SELECT * FROM financial_ratios",
        conn,
    )

    pnl = pd.read_sql_query(
        """
        SELECT company_id, year, operating_profit
        FROM profitandloss
        """,
        conn,
    )

    bs = pd.read_sql_query(
        """
        SELECT company_id, year, equity_capital, reserves, borrowings
        FROM balancesheet
        """,
        conn,
    )

    companies = pd.read_sql_query(
        "SELECT * FROM companies",
        conn,
    )

    conn.close()

    peer_groups = pd.read_excel(PEER_GROUPS_PATH)

    return ratios, pnl, bs, companies, peer_groups


def calculate_roce(ratios, pnl, bs):
    pnl = pnl.copy()
    bs = bs.copy()

    pnl["year"] = pd.to_numeric(pnl["year"], errors="coerce")
    bs["year"] = pd.to_numeric(bs["year"], errors="coerce")
    ratios["year"] = pd.to_numeric(ratios["year"], errors="coerce")

    merged = pnl.merge(
        bs,
        on=["company_id", "year"],
        how="left",
    )

    merged["capital_employed"] = (
        merged["equity_capital"]
        + merged["reserves"]
        + merged["borrowings"]
    )

    merged["roce_pct"] = (
        merged["operating_profit"]
        / merged["capital_employed"]
        * 100
    )

    roce = merged[
        ["company_id", "year", "roce_pct"]
    ].copy()

    return ratios.merge(
        roce,
        on=["company_id", "year"],
        how="left",
    )


def latest_ratios(ratios):
    ratios = ratios.sort_values(
        ["company_id", "year"]
    )

    return (
        ratios
        .groupby("company_id", as_index=False)
        .tail(1)
        .copy()
    )


def calculate_percentiles(df):
    result = df.copy()

    for metric in METRICS:
        source = METRIC_SOURCE[metric]

        if source not in result.columns:
            result[source] = pd.NA

        values = pd.to_numeric(
            result[source],
            errors="coerce",
        )

        if metric == "D/E":
            result[f"{metric}_Percentile"] = (
                1
                - values.rank(
                    method="min",
                    pct=True,
                )
            ) * 100
        else:
            result[f"{metric}_Percentile"] = (
                values.rank(
                    method="min",
                    pct=True,
                )
            ) * 100

    return result


def get_company_name(companies):
    return companies[
        ["id", "company_name"]
    ].rename(
        columns={
            "id": "company_id"
        }
    )

    if "name" in columns:
        return companies[
            ["company_id", columns["name"]]
        ].rename(
            columns={
                columns["name"]: "company_name"
            }
        )

    return companies[["company_id"]].assign(
        company_name=companies["company_id"]
    )


def build_output():
    ratios, pnl, bs, companies, peer_groups = load_data()

    ratios = calculate_roce(
        ratios,
        pnl,
        bs,
    )

    ratios = latest_ratios(ratios)

    company_names = get_company_name(companies)

    peer_groups = peer_groups[
        [
            "peer_group_name",
            "company_id",
            "is_benchmark",
        ]
    ].copy()

    peer_groups["company_id"] = (
        peer_groups["company_id"]
        .astype(str)
        .str.strip()
    )

    ratios["company_id"] = (
        ratios["company_id"]
        .astype(str)
        .str.strip()
    )

    data = peer_groups.merge(
        ratios,
        on="company_id",
        how="left",
    )

    data = data.merge(
        company_names,
        on="company_id",
        how="left",
    )

    data = calculate_percentiles(data)

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        OUTPUT_PATH,
        engine="openpyxl",
    ) as writer:

        peer_names = (
            peer_groups["peer_group_name"]
            .dropna()
            .drop_duplicates()
            .tolist()
        )

        for peer_name in peer_names:
            group = data[
                data["peer_group_name"] == peer_name
            ].copy()

            columns = [
                "company_id",
                "company_name",
            ]

            for metric in METRICS:
                source = METRIC_SOURCE[metric]
                percentile = f"{metric}_Percentile"

                columns.append(source)
                columns.append(percentile)

            columns.append("is_benchmark")

            export = group[
                [
                    c
                    for c in columns
                    if c in group.columns
                ]
            ].copy()

            export = export.sort_values(
                "is_benchmark",
                ascending=False,
            )

            export.to_excel(
                writer,
                sheet_name=peer_name[:31],
                index=False,
            )

    apply_formatting()

    print("Day 20 completed")
    print(f"Output: {OUTPUT_PATH}")
    print(f"Sheets generated: {len(peer_names)}")


def apply_formatting():
    workbook = load_workbook(OUTPUT_PATH)

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFEB9C",
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    benchmark_fill = PatternFill(
        fill_type="solid",
        fgColor="FFD966",
    )

    for worksheet in workbook.worksheets:
        headers = {
            cell.value: cell.column
            for cell in worksheet[1]
        }

        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        for metric in METRICS:
            percentile_column = f"{metric}_Percentile"

            if percentile_column not in headers:
                continue

            column_number = headers[
                percentile_column
            ]

            column_letter = get_column_letter(
                column_number
            )

            cell_range = (
                f"{column_letter}2:"
                f"{column_letter}{worksheet.max_row}"
            )

            worksheet.conditional_formatting.add(
                cell_range,
                CellIsRule(
                    operator="greaterThanOrEqual",
                    formula=["75"],
                    fill=green_fill,
                ),
            )

            worksheet.conditional_formatting.add(
                cell_range,
                CellIsRule(
                    operator="between",
                    formula=["25", "74.999999"],
                    fill=yellow_fill,
                ),
            )

            worksheet.conditional_formatting.add(
                cell_range,
                CellIsRule(
                    operator="lessThan",
                    formula=["25"],
                    fill=red_fill,
                ),
            )

        benchmark_column = headers.get(
            "is_benchmark"
        )

        if benchmark_column:
            for row in range(
                2,
                worksheet.max_row + 1,
            ):
                value = worksheet.cell(
                    row=row,
                    column=benchmark_column,
                ).value

                if value in (
                    1,
                    True,
                    "1",
                    "TRUE",
                    "True",
                ):
                    for cell in worksheet[row]:
                        cell.fill = benchmark_fill

        metric_columns = [
            headers.get(METRIC_SOURCE[m])
            for m in METRICS
        ]

        for row in range(
            2,
            worksheet.max_row + 1,
        ):
            for column in metric_columns:
                if column:
                    worksheet.cell(
                        row=row,
                        column=column,
                    ).number_format = "0.00"

        summary_row = worksheet.max_row + 2

        worksheet.cell(
            summary_row,
            1,
            "Peer Group Median",
        ).font = Font(bold=True)

        for metric in METRICS:
            source = METRIC_SOURCE[metric]

            if source not in headers:
                continue

            column = headers[source]

            values = []

            for row in range(
                2,
                worksheet.max_row + 1,
            ):
                value = worksheet.cell(
                    row=row,
                    column=column,
                ).value

                if isinstance(
                    value,
                    (int, float),
                ):
                    values.append(value)

            if values:
                worksheet.cell(
                    summary_row,
                    column,
                    pd.Series(values).median(),
                )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(worksheet.max_column)}"
            f"{worksheet.max_row}"
        )

        for column_cells in worksheet.columns:
            maximum = 0

            for cell in column_cells:
                if cell.value is not None:
                    maximum = max(
                        maximum,
                        len(str(cell.value)),
                    )

            column_letter = get_column_letter(
                column_cells[0].column
            )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(maximum + 2, 12),
                30,
            )

    workbook.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_output()